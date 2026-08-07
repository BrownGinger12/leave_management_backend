from pydantic import BaseModel  # import BaseModel as the base for all models
from gateway.mysql_gateway import fetch_query  # import fetch_query for SELECT operations


class LeaveTransaction(BaseModel):
    """
    Model for querying and exporting leave transaction records.
    A leave transaction is one leave application with all its incurred dates
    aggregated into a single row for reporting purposes.
    """

    @staticmethod
    def export_to_excel(date_from: str, date_to: str):
        """
        Generates an Excel report of all leave applications whose leave dates
        fall within the given date range using the official DepEd template
        (Template/LEAVE TRANSACTIONS.xlsx).
        Each row represents one leave application. Column K lists all
        individual leave dates incurred (comma-separated).

        Parameters:
            date_from (str): Start of the leave date range (YYYY-MM-DD).
            date_to (str): End of the leave date range (YYYY-MM-DD).

        Returns:
            io.BytesIO: In-memory Excel file ready to be sent as an attachment.
            dict: Error dict if a problem occurs before the file is generated.
        """
        import os  # used to build the template path
        import io  # used to create an in-memory byte buffer
        import openpyxl  # used to load and write the Excel workbook
        from openpyxl.styles import Border, Side, Font, Alignment  # cell styling classes

        try:
            rows = fetch_query(  # fetch one row per application; leave dates aggregated via GROUP_CONCAT
                """SELECT
                        e.leave_card_number,
                        la.application_number,
                        e.first_name, e.last_name, e.middle_name,
                        e.sex,
                        e.position,
                        e.employee_type,
                        s.name AS school_name,
                        la.date_filed,
                        la.remarks,
                        la.other_leave_description,
                        lt.name AS leave_type_name,
                        GROUP_CONCAT(DATE_FORMAT(lad.leave_date, '%%m/%%d/%%Y')
                                     ORDER BY lad.leave_date ASC SEPARATOR ', ') AS leave_dates,
                        SUM(CASE WHEN lad.duration_type = 'HALF_DAY' THEN 0.5 ELSE 1.0 END) AS total_days
                    FROM leave_applications la
                    JOIN leave_application_dates lad ON lad.leave_application_id = la.id
                    JOIN employees e ON e.id = la.employee_id
                    JOIN leave_types lt ON lt.id = la.leave_type_id
                    LEFT JOIN schools s ON s.id = e.school_id
                    WHERE la.status NOT IN ('RETURNED', 'DISAPPROVED')
                      AND la.is_deleted = 0
                      AND e.is_active = 1
                      AND lad.leave_date BETWEEN %s AND %s
                    GROUP BY la.id, e.id, e.leave_card_number, la.application_number,
                             e.first_name, e.last_name, e.middle_name, e.sex,
                             e.position, e.employee_type, s.name,
                             la.date_filed, la.remarks, la.other_leave_description,
                             lt.name
                    ORDER BY MIN(lad.leave_date) ASC, e.last_name ASC, e.first_name ASC""",
                [date_from, date_to]
            ) or []

            template_path = os.path.join(  # resolve absolute path to the leave transactions template
                os.path.dirname(os.path.dirname(__file__)), "Template", "LEAVE TRANSACTIONS.xlsx"
            )
            wb = openpyxl.load_workbook(template_path)  # load template (preserves header image and formatting)
            ws = wb.active  # use the first (and only) sheet

            ws.unmerge_cells("A50:M50")  # unmerge page-counter row so it can be relocated

            thin = Side(style="thin")  # thin border side definition
            thin_border = Border(left=thin, right=thin, top=thin, bottom=thin)  # all-sides thin border

            DATA_START = 15  # row where data rows begin in this template

            def fmt_date(d):
                """Format a date object or None as MM/DD/YYYY string."""
                if d is None:  # no date available
                    return ""  # return empty string for missing dates
                return d.strftime("%m/%d/%Y") if hasattr(d, "strftime") else str(d)  # format MM/DD/YYYY

            SEX_MAP = {"MALE": "M", "FEMALE": "F"}  # map full sex value to single-letter abbreviation

            TYPE_MAP = {  # map database employee_type to display label
                "TEACHING": "Teaching",
                "NON_TEACHING": "Non-Teaching",
            }

            for i, row in enumerate(rows):  # write one row per leave application
                r = DATA_START + i  # absolute sheet row for this application

                last = (row["last_name"] or "").strip()  # employee last name
                first = (row["first_name"] or "").strip()  # employee first name
                middle = (row.get("middle_name") or "").strip()  # middle name (may be None)
                mi = f" {middle[0]}." if middle else ""  # middle initial with period, or empty string
                full_name = f"{last}, {first}{mi}"  # format: LAST, FIRST MI.

                sex = SEX_MAP.get((row.get("sex") or "").upper(), "")  # M or F, blank if unknown
                emp_type = TYPE_MAP.get(row.get("employee_type") or "", row.get("employee_type") or "")  # readable type

                leave_type = row["leave_type_name"]  # full leave type name (e.g., "Sick Leave")

                remarks = row.get("remarks") or ""  # leave application remarks
                if row.get("other_leave_description"):  # for "Others" leave type, use the specific description
                    remarks = row["other_leave_description"]  # override with the specific description

                values = [  # ordered by column A–M (13 columns)
                    i + 1,                                    # A: sequential row number
                    row.get("leave_card_number") or "",       # B: leave card number
                    row.get("application_number") or "",      # C: application number
                    full_name,                                # D: formatted employee name
                    sex,                                      # E: M or F
                    row.get("position") or "",                # F: position/designation
                    emp_type,                                 # G: Teaching or Non-Teaching
                    row.get("school_name") or "",             # H: school or division name
                    fmt_date(row.get("date_filed")),          # I: date application was filed
                    leave_type,                               # J: type of leave (full name)
                    row.get("leave_dates") or "",             # K: all leave dates incurred (comma-separated)
                    float(row["total_days"] or 0),            # L: total number of leave days
                    remarks,                                  # M: remarks / reason
                ]

                for col_idx, val in enumerate(values, start=1):  # write each value to its column cell
                    cell = ws.cell(row=r, column=col_idx)  # target cell
                    cell.value = val  # set value
                    cell.border = thin_border  # apply thin border to match header style

            for old_row in [46, 48, 50]:  # clear old static footer rows from template
                ws.cell(row=old_row, column=1).value = None  # clear only master cell (column A)
            ws.cell(row=48, column=2).value = None  # clear signature line in column B

            footer_row = max(46, DATA_START + len(rows) + 3)  # footer at least at row 46, or below data

            cert_cell = ws.cell(row=footer_row, column=1)  # CERTIFIED CORRECT label cell
            cert_cell.value = "CERTIFIED CORRECT:"  # footer label
            cert_cell.font = Font(bold=True)  # bold to match template style

            ws.cell(row=footer_row + 2, column=2).value = "_________________________"  # signature line

            page_row = footer_row + 4  # row for the page counter
            ws.merge_cells(f"A{page_row}:M{page_row}")  # merge across all 13 columns
            page_cell = ws.cell(row=page_row, column=1)  # write to the master cell of the merge
            page_cell.value = "page 1 of 1"  # static page label
            page_cell.alignment = Alignment(horizontal="center")  # centre-align the label

            buffer = io.BytesIO()  # create in-memory byte buffer
            wb.save(buffer)  # write workbook to buffer
            buffer.seek(0)  # rewind to beginning before returning
            return buffer  # caller streams this as a file download

        except Exception as e:  # catch unexpected errors (missing template, DB failure, etc.)
            return {"statusCode": 500, "message": str(e)}
