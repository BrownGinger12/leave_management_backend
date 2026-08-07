from pydantic import BaseModel  # import BaseModel as the base for all models
from gateway.mysql_gateway import fetch_query  # import fetch_query for SELECT operations
from datetime import date  # import date for default year range


class LeaveWithoutPay(BaseModel):
    """
    Model for querying leave without pay (LWOP) records.
    LWOP dates are leave_application_dates rows where is_paid = 0.
    Excludes RETURNED and DISAPPROVED applications.
    """

    @staticmethod
    def get_paginated(employee_type: str, date_from: str, date_to: str,
                      page: int = 1, limit: int = 10,
                      school_type: str = None) -> dict:
        """
        Returns a paginated list of leave-without-pay dates for employees of
        the given type within the specified date range.
        Optionally filters by school_type (REGULAR, IU, or SHS).
        Each row represents one LWOP date tied to a leave application.

        Parameters:
            employee_type (str): 'TEACHING' or 'NON_TEACHING'.
            date_from (str): Start of the date range (YYYY-MM-DD).
            date_to (str): End of the date range (YYYY-MM-DD).
            page (int): Page number (1-indexed, default 1).
            limit (int): Records per page (default 10).
            school_type (str): Optional — filter by school classification ('REGULAR', 'IU', 'SHS').

        Returns:
            dict: statusCode 200 with paginated LWOP records, or an error dict.
        """
        try:
            page = max(page, 1)  # ensure page is at least 1
            limit = max(limit, 1)  # ensure limit is at least 1
            offset = (page - 1) * limit  # compute SQL offset from page number

            valid_school_types = ("REGULAR", "IU", "SHS")  # accepted school_type values
            if school_type and school_type.upper() not in valid_school_types:  # validate if provided
                return {"statusCode": 400, "message": f"school_type must be one of: {', '.join(valid_school_types)}"}

            school_type_filter = ""  # additional filter fragment; empty when not filtering by school_type
            params = [employee_type, date_from, date_to]  # base query parameters

            if school_type:  # append school_type condition when provided
                school_type_filter = "AND s.school_type = %s"  # filter on the schools table
                params.append(school_type.upper())  # bind the normalised value

            base_where = f"""
                FROM leave_application_dates lad
                JOIN leave_applications la ON la.id = lad.leave_application_id
                JOIN employees e ON e.id = la.employee_id
                JOIN leave_types lt ON lt.id = la.leave_type_id
                LEFT JOIN schools s ON s.id = e.school_id
                WHERE lad.is_paid = 0
                  AND la.status NOT IN ('RETURNED', 'DISAPPROVED')
                  AND la.is_deleted = 0
                  AND e.is_active = 1
                  AND e.employee_type = %s
                  AND lad.leave_date BETWEEN %s AND %s
                  {school_type_filter}
            """  # shared WHERE clause reused for both count and data queries

            total_row = fetch_query(  # count total matching LWOP date rows
                f"SELECT COUNT(*) AS total {base_where}",
                params
            )
            total = int(total_row[0]["total"]) if total_row else 0  # cast to int

            rows = fetch_query(  # fetch paginated LWOP records with school_type included
                f"""SELECT e.id AS employee_id,
                           e.first_name, e.last_name, e.employee_number,
                           e.employee_type, e.position,
                           e.implementing_unit,
                           s.name AS school_name,
                           s.school_type,
                           la.id AS application_id,
                           la.application_number,
                           la.status,
                           la.date_filed,
                           lt.code AS leave_type_code,
                           lt.name AS leave_type_name,
                           lad.id AS leave_date_id,
                           lad.leave_date,
                           lad.duration_type,
                           lad.half_day_period,
                           CASE WHEN lad.duration_type = 'HALF_DAY' THEN 0.5 ELSE 1.0 END AS days_without_pay
                    {base_where}
                    ORDER BY lad.leave_date DESC, e.last_name ASC, e.first_name ASC
                    LIMIT %s OFFSET %s""",
                params + [limit, offset]
            )

            total_pages = (total + limit - 1) // limit if total > 0 else 1  # ceiling division for total pages

            return {  # return paginated response
                "statusCode": 200,  # success code
                "employee_type": employee_type,  # employee type filter applied
                "school_type": school_type.upper() if school_type else None,  # school type filter applied (or None)
                "date_from": date_from,  # range start
                "date_to": date_to,  # range end
                "page": page,  # current page number
                "limit": limit,  # records per page
                "total": total,  # total matching LWOP date rows
                "total_pages": total_pages,  # total number of pages
                "count": len(rows) if rows else 0,  # records on this page
                "data": rows if rows else [],  # LWOP date records
            }

        except Exception as e:  # catch unexpected errors
            return {"statusCode": 500, "message": str(e)}  # return 500 with error detail

    # --------------------------
    # Export teaching LWOP to Excel
    # --------------------------

    @staticmethod
    def export_teaching_to_excel(date_from: str, date_to: str, school_type: str = None):
        """
        Generates an Excel report of TEACHING leave-without-pay applications for the
        given date range using the official DepEd template
        (Template/LEAVE WITHOUT PAY TEACHING.xlsx).
        Each row represents one leave application grouped with aggregated paid/unpaid dates.

        Parameters:
            date_from (str): Start of the LWOP date range (YYYY-MM-DD).
            date_to (str): End of the LWOP date range (YYYY-MM-DD).
            school_type (str): Optional — filter by school classification (REGULAR, IU, SHS).

        Returns:
            io.BytesIO: In-memory Excel file ready to be sent as an attachment.
            dict: Error dict if a problem occurs before the file is generated.
        """
        import os  # used to build the template path
        import io  # used to create an in-memory byte buffer
        import openpyxl  # used to load and write the Excel workbook
        from openpyxl.styles import Border, Side, Font, Alignment  # cell styling classes

        try:
            school_type_filter = ""  # extra WHERE fragment for school_type
            params = []  # outer-query params (school_type only, if provided)

            if school_type:  # append school_type condition when provided
                school_type_filter = "AND s.school_type = %s"  # filter on schools table
                params.append(school_type.upper())  # bind normalised value

            rows = fetch_query(  # fetch one row per application with comma-separated date lists
                f"""SELECT
                        e.leave_card_number,
                        e.first_name, e.last_name, e.middle_name,
                        e.salary, e.employee_number,
                        s.name AS school_name,
                        la.date_filed, la.status,
                        la.remarks,
                        la.other_leave_description,
                        lt.code AS leave_type_code,
                        lt.name AS leave_type_name,
                        GROUP_CONCAT(CASE WHEN lad.is_paid = 1 THEN DATE_FORMAT(lad.leave_date, '%%m/%%d/%%Y') END
                                     ORDER BY lad.leave_date ASC SEPARATOR ', ') AS paid_dates,
                        SUM(CASE WHEN lad.is_paid = 1 AND lad.duration_type = 'HALF_DAY' THEN 0.5
                                 WHEN lad.is_paid = 1 THEN 1.0 ELSE 0 END) AS days_with_pay,
                        GROUP_CONCAT(CASE WHEN lad.is_paid = 0 THEN DATE_FORMAT(lad.leave_date, '%%m/%%d/%%Y') END
                                     ORDER BY lad.leave_date ASC SEPARATOR ', ') AS unpaid_dates,
                        SUM(CASE WHEN lad.is_paid = 0 AND lad.duration_type = 'HALF_DAY' THEN 0.5
                                 WHEN lad.is_paid = 0 THEN 1.0 ELSE 0 END) AS days_without_pay,
                        GROUP_CONCAT(DATE_FORMAT(lad.leave_date, '%%m/%%d/%%Y')
                                     ORDER BY lad.leave_date ASC SEPARATOR ', ') AS total_dates,
                        SUM(CASE WHEN lad.duration_type = 'HALF_DAY' THEN 0.5 ELSE 1.0 END) AS total_days
                    FROM leave_applications la
                    JOIN leave_application_dates lad ON lad.leave_application_id = la.id
                    JOIN employees e ON e.id = la.employee_id
                    JOIN leave_types lt ON lt.id = la.leave_type_id
                    LEFT JOIN schools s ON s.id = e.school_id
                    WHERE la.status NOT IN ('RETURNED', 'DISAPPROVED')
                      AND la.is_deleted = 0
                      AND e.is_active = 1
                      AND e.employee_type = 'TEACHING'
                      {school_type_filter}
                      AND la.id IN (
                          SELECT DISTINCT leave_application_id
                          FROM leave_application_dates
                          WHERE is_paid = 0 AND leave_date BETWEEN %s AND %s
                      )
                    GROUP BY la.id, e.id, e.leave_card_number, e.first_name, e.last_name,
                             e.middle_name, e.salary, e.employee_number,
                             s.name, la.date_filed, la.status,
                             la.remarks, lt.code, lt.name, la.other_leave_description
                    ORDER BY MIN(lad.leave_date) ASC, e.last_name ASC, e.first_name ASC""",
                params + [date_from, date_to]  # school_type param first, then date range for subquery
            ) or []

            template_path = os.path.join(  # resolve absolute path to the LWOP teaching template
                os.path.dirname(os.path.dirname(__file__)), "Template", "LEAVE WITHOUT PAY TEACHING.xlsx"
            )
            wb = openpyxl.load_workbook(template_path)  # load template (preserves header image and formatting)
            ws = wb.active  # use the first sheet

            ws.unmerge_cells("A59:P59")  # unmerge page-counter row so we can relocate it freely

            thin = Side(style="thin")  # thin border side
            thin_border = Border(left=thin, right=thin, top=thin, bottom=thin)  # all-sides thin border

            DATA_START = 16  # row where data rows begin in the template

            def fmt_date(d):
                """Format a date object or None as MM/DD/YYYY string."""
                if d is None:  # no date available for this column
                    return ""  # return empty string
                return d.strftime("%m/%d/%Y") if hasattr(d, "strftime") else str(d)  # format as MM/DD/YYYY

            for i, row in enumerate(rows):  # write one row per leave application
                r = DATA_START + i  # absolute sheet row number

                last = (row["last_name"] or "").strip()  # employee last name
                first = (row["first_name"] or "").strip()  # employee first name
                middle = (row.get("middle_name") or "").strip()  # middle name (may be None)
                mi = f" {middle[0]}." if middle else ""  # middle initial with period, or empty
                full_name = f"{last}, {first}{mi}"  # format: LAST, FIRST MI.

                leave_type = row["leave_type_code"]  # leave type code (VL, SL, etc.)
                if leave_type == "OT" and row.get("other_leave_description"):  # "Others" — append description
                    leave_type = f"Others ({row['other_leave_description']})"  # show specific leave

                salary = float(row["salary"]) if row.get("salary") else ""  # cast Decimal to float or blank

                values = [  # ordered by column A–P
                    i + 1,                                         # A: sequential row number
                    row.get("leave_card_number") or "",            # B: leave card number
                    full_name,                                     # C: formatted name
                    salary,                                        # D: salary
                    row.get("employee_number") or "",              # E: employee number
                    row.get("school_name") or "",                  # F: school/division
                    fmt_date(row.get("date_filed")),               # G: date filed
                    leave_type,                                    # H: type of leave
                    row.get("paid_dates") or "",                   # I: all paid dates (comma-separated)
                    row.get("unpaid_dates") or "",                 # J: all unpaid dates (comma-separated)
                    row.get("total_dates") or "",                  # K: all leave dates incurred (comma-separated)
                    float(row["days_with_pay"] or 0),              # L: days with pay
                    float(row["days_without_pay"] or 0),                         # M: days without pay
                    float(row["total_days"] or 0),                               # N: total days
                    row.get("remarks") or row.get("other_leave_description") or "", # O: remarks / reason
                    "",                                                            # P: check if reinstated (manual)
                ]

                for col_idx, val in enumerate(values, start=1):  # write each column cell
                    cell = ws.cell(row=r, column=col_idx)  # target cell
                    cell.value = val  # set cell value
                    cell.border = thin_border  # apply thin border to match header style

            for old_row in [55, 57, 59]:  # clear old static footer rows from template
                ws.cell(row=old_row, column=1).value = None  # only clear master cell (column A)

            footer_row = max(55, DATA_START + len(rows) + 3)  # footer at least at row 55, or below data

            cert_cell = ws.cell(row=footer_row, column=1)  # CERTIFIED CORRECT label cell
            cert_cell.value = "CERTIFIED CORRECT:"  # footer label
            cert_cell.font = Font(bold=True)  # bold to match template style

            ws.cell(row=footer_row + 2, column=2).value = "_________________________"  # signature line

            page_row = footer_row + 4  # row for the page counter
            ws.merge_cells(f"A{page_row}:P{page_row}")  # merge full row for page counter (16 columns)
            page_cell = ws.cell(row=page_row, column=1)  # write to master cell of the merge
            page_cell.value = "page 1 of 1"  # static page label
            page_cell.alignment = Alignment(horizontal="center")  # centre-align

            buffer = io.BytesIO()  # create in-memory byte buffer
            wb.save(buffer)  # write workbook to buffer
            buffer.seek(0)  # rewind to start before sending
            return buffer  # caller streams this as a file download

        except Exception as e:  # catch unexpected errors (missing template, DB failure, etc.)
            return {"statusCode": 500, "message": str(e)}

    # --------------------------
    # Export non-teaching LWOP to Excel
    # --------------------------

    @staticmethod
    def export_non_teaching_to_excel(date_from: str, date_to: str, school_type: str = None):
        """
        Generates an Excel report of NON_TEACHING leave-without-pay applications for the
        given date range using the official DepEd template
        (Template/LEAVE WITHOUT PAY NON TEACHING.xlsx).
        Each row represents one leave application. Unpaid days are split into VL (col M)
        and SL (col N) to match the non-teaching template layout.

        Parameters:
            date_from (str): Start of the LWOP date range (YYYY-MM-DD).
            date_to (str): End of the LWOP date range (YYYY-MM-DD).
            school_type (str): Optional — filter by school classification (REGULAR, IU, SHS).

        Returns:
            io.BytesIO: In-memory Excel file ready to be sent as an attachment.
            dict: Error dict if a problem occurs before the file is generated.
        """
        import os  # used to build the template path
        import io  # used to create an in-memory byte buffer
        import openpyxl  # used to load and write the Excel workbook
        from openpyxl.styles import Border, Side, Font, Alignment  # cell styling classes

        try:
            school_type_filter = ""  # extra WHERE fragment for school_type
            params = []  # outer-query params (school_type only, if provided)

            if school_type:  # append school_type condition when provided
                school_type_filter = "AND s.school_type = %s"  # filter on schools table
                params.append(school_type.upper())  # bind normalised value

            rows = fetch_query(  # fetch one row per application with comma-separated date lists and day totals
                f"""SELECT
                        e.leave_card_number,
                        e.first_name, e.last_name, e.middle_name,
                        e.salary, e.employee_number,
                        s.name AS school_name,
                        la.date_filed, la.status,
                        la.remarks,
                        la.other_leave_description,
                        lt.code AS leave_type_code,
                        lt.name AS leave_type_name,
                        GROUP_CONCAT(CASE WHEN lad.is_paid = 1 THEN DATE_FORMAT(lad.leave_date, '%%m/%%d/%%Y') END
                                     ORDER BY lad.leave_date ASC SEPARATOR ', ') AS paid_dates,
                        SUM(CASE WHEN lad.is_paid = 1 AND lad.duration_type = 'HALF_DAY' THEN 0.5
                                 WHEN lad.is_paid = 1 THEN 1.0 ELSE 0 END) AS days_with_pay,
                        GROUP_CONCAT(CASE WHEN lad.is_paid = 0 THEN DATE_FORMAT(lad.leave_date, '%%m/%%d/%%Y') END
                                     ORDER BY lad.leave_date ASC SEPARATOR ', ') AS unpaid_dates,
                        SUM(CASE WHEN lad.is_paid = 0 AND lad.duration_type = 'HALF_DAY' THEN 0.5
                                 WHEN lad.is_paid = 0 THEN 1.0 ELSE 0 END) AS days_without_pay,
                        GROUP_CONCAT(DATE_FORMAT(lad.leave_date, '%%m/%%d/%%Y')
                                     ORDER BY lad.leave_date ASC SEPARATOR ', ') AS total_dates,
                        SUM(CASE WHEN lad.duration_type = 'HALF_DAY' THEN 0.5 ELSE 1.0 END) AS total_days
                    FROM leave_applications la
                    JOIN leave_application_dates lad ON lad.leave_application_id = la.id
                    JOIN employees e ON e.id = la.employee_id
                    JOIN leave_types lt ON lt.id = la.leave_type_id
                    LEFT JOIN schools s ON s.id = e.school_id
                    WHERE la.status NOT IN ('RETURNED', 'DISAPPROVED')
                      AND la.is_deleted = 0
                      AND e.is_active = 1
                      AND e.employee_type = 'NON_TEACHING'
                      {school_type_filter}
                      AND la.id IN (
                          SELECT DISTINCT leave_application_id
                          FROM leave_application_dates
                          WHERE is_paid = 0 AND leave_date BETWEEN %s AND %s
                      )
                    GROUP BY la.id, e.id, e.leave_card_number, e.first_name, e.last_name,
                             e.middle_name, e.salary, e.employee_number,
                             s.name, la.date_filed, la.status,
                             la.remarks, lt.code, lt.name, la.other_leave_description
                    ORDER BY MIN(lad.leave_date) ASC, e.last_name ASC, e.first_name ASC""",
                params + [date_from, date_to]
            ) or []

            template_path = os.path.join(  # resolve absolute path to the non-teaching LWOP template
                os.path.dirname(os.path.dirname(__file__)), "Template", "LEAVE WITHOUT PAY NON TEACHING.xlsx"
            )
            wb = openpyxl.load_workbook(template_path)  # load template (preserves header image and formatting)
            ws = wb.active  # use the first sheet

            ws.unmerge_cells("A65:Q65")  # unmerge page-counter row so it can be relocated

            thin = Side(style="thin")  # thin border side
            thin_border = Border(left=thin, right=thin, top=thin, bottom=thin)  # all-sides thin border

            DATA_START = 18  # row where data rows begin in the non-teaching template

            def fmt_date(d):
                """Format a date object or None as MM/DD/YYYY string."""
                if d is None:  # no date available
                    return ""  # empty string for missing dates
                return d.strftime("%m/%d/%Y") if hasattr(d, "strftime") else str(d)  # format MM/DD/YYYY

            VL_CODES = {"VL", "FL"}  # leave types whose unpaid days count as VL without pay

            for i, row in enumerate(rows):  # write one row per leave application
                r = DATA_START + i  # absolute sheet row number

                last = (row["last_name"] or "").strip()  # employee last name
                first = (row["first_name"] or "").strip()  # employee first name
                middle = (row.get("middle_name") or "").strip()  # middle name (may be None)
                mi = f" {middle[0]}." if middle else ""  # middle initial with period, or empty
                full_name = f"{last}, {first}{mi}"  # format: LAST, FIRST MI.

                leave_type = row["leave_type_code"]  # leave type code (VL, SL, FL, etc.)
                if leave_type == "OT" and row.get("other_leave_description"):  # Others — append description
                    leave_type = f"Others ({row['other_leave_description']})"  # show specific type

                salary = float(row["salary"]) if row.get("salary") else ""  # cast Decimal to float or blank
                days_wo_pay = float(row["days_without_pay"] or 0)  # total unpaid days for this application

                # Split unpaid days into VL column (M) and SL column (N)
                vl_wo_pay = days_wo_pay if row["leave_type_code"] in VL_CODES else 0.0  # VL/FL unpaid days
                sl_wo_pay = days_wo_pay if row["leave_type_code"] == "SL" else 0.0  # SL unpaid days

                values = [  # ordered by column A–Q (17 columns)
                    i + 1,                                     # A: sequential row number
                    row.get("leave_card_number") or "",        # B: leave card number
                    full_name,                                 # C: formatted name
                    salary,                                    # D: salary
                    row.get("employee_number") or "",          # E: employee number
                    row.get("school_name") or "",              # F: school/division
                    fmt_date(row.get("date_filed")),           # G: date filed
                    leave_type,                                # H: type of leave
                    row.get("paid_dates") or "",               # I: all paid dates (comma-separated)
                    row.get("unpaid_dates") or "",             # J: all unpaid dates (comma-separated)
                    row.get("total_dates") or "",              # K: all leave dates incurred (comma-separated)
                    float(row["days_with_pay"] or 0),          # L: days with pay
                    vl_wo_pay,                                 # M: VL days without pay
                    sl_wo_pay,                                 # N: SL days without pay
                    float(row["total_days"] or 0),                               # O: total days of leave
                    row.get("remarks") or row.get("other_leave_description") or "", # P: remarks / reason
                    "",                                                            # Q: check if reinstated (manual)
                ]

                for col_idx, val in enumerate(values, start=1):  # write each column cell
                    cell = ws.cell(row=r, column=col_idx)  # target cell
                    cell.value = val  # set cell value
                    cell.border = thin_border  # apply thin border to match header style

            for old_row in [61, 63, 65]:  # clear old static footer rows from template
                ws.cell(row=old_row, column=1).value = None  # only clear master cell (column A)

            footer_row = max(61, DATA_START + len(rows) + 3)  # footer at least at row 61, or below data

            cert_cell = ws.cell(row=footer_row, column=1)  # CERTIFIED CORRECT label cell
            cert_cell.value = "CERTIFIED CORRECT:"  # footer label
            cert_cell.font = Font(bold=True)  # bold to match template style

            ws.cell(row=footer_row + 2, column=2).value = "_________________________"  # signature line

            page_row = footer_row + 4  # row for the page counter
            ws.merge_cells(f"A{page_row}:Q{page_row}")  # merge full row (17 columns) for page counter
            page_cell = ws.cell(row=page_row, column=1)  # write to master cell of the merge
            page_cell.value = "page 1 of 1"  # static page label
            page_cell.alignment = Alignment(horizontal="center")  # centre-align

            buffer = io.BytesIO()  # create in-memory byte buffer
            wb.save(buffer)  # write workbook to buffer
            buffer.seek(0)  # rewind to start before sending
            return buffer  # caller streams this as a file download

        except Exception as e:  # catch unexpected errors
            return {"statusCode": 500, "message": str(e)}
